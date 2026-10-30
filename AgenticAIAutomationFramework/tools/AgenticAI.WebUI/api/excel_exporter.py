# Exports SQA-style test cases (see llm_test_generator.py) to an .xlsx workbook,
# one row per test case, ready to paste into TestRail/Jira/Excel-based QA trackers.

import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ('id', 'Test Case ID', 16),
    ('title', 'Title', 42),
    ('module', 'Module', 16),
    ('priority', 'Priority', 10),
    ('type', 'Type', 10),
    ('preconditions', 'Preconditions', 36),
    ('steps_text', 'Test Steps', 50),
    ('test_data', 'Test Data', 20),
    ('expected', 'Expected Result', 36),
    ('postconditions', 'Postconditions', 36),
    ('tags_text', 'Tags', 20),
]

HEADER_FILL = PatternFill(start_color='0F766E', end_color='0F766E', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')


def _steps_text(test_case):
    detailed = test_case.get('detailed_steps')
    if detailed:
        return '\n'.join(f"{s['step_number']}. {s['action']} -> {s['expected']}" for s in detailed)
    steps = test_case.get('steps') or []
    return '\n'.join(f"{i + 1}. {s}" for i, s in enumerate(steps))


def export_test_cases_to_excel(test_cases: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Test Cases'

    for col_index, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.freeze_panes = 'A2'

    for row_index, tc in enumerate(test_cases, start=2):
        row_values = {
            'id': tc.get('id', ''),
            'title': tc.get('title') or tc.get('name', ''),
            'module': tc.get('module', ''),
            'priority': tc.get('priority', ''),
            'type': tc.get('type', ''),
            'preconditions': tc.get('preconditions', ''),
            'steps_text': _steps_text(tc),
            'test_data': tc.get('test_data', ''),
            'expected': tc.get('expected', ''),
            'postconditions': tc.get('postconditions', ''),
            'tags_text': ', '.join(tc.get('tags') or []),
        }
        for col_index, (key, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=row_values[key])
            cell.alignment = WRAP

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
